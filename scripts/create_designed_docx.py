"""Editable designed resume (.docx) built to survive import into Google Docs.

Google Docs drops floating shapes and picture crop shapes, so this version uses
only features Docs imports reliably: a two-column table whose left cell is
shaded navy, a photo pre-cut into a circle (transparent PNG), paragraph
borders, standard bullets, and Google fonts.

Edit content in resume.md, then run build.command.
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw

from create_designed_resume import PHOTO, split_lead
from resume_content import (
    CERTIFICATIONS, CONTACT, CORE_COMPETENCIES, EDUCATION, EXPERIENCE, HEADLINE, HIGHLIGHTS, NAME,
    SIDEBAR_SKILLS, SUMMARY, TAGLINE,
)

ROOT = Path(__file__).resolve().parent.parent  # the resume folder (scripts/ lives inside it)
DOCX_OUT = ROOT / "Abhishek_Gupta_Resume_Designed.docx"
ROUND_PHOTO = ROOT / "assets" / "photo-round.png"

FONT = "Lato"  # available in Google Docs
NAVY_HEX = "1B3A5C"
NAVY = RGBColor(0x1B, 0x3A, 0x5C)
ACCENT = RGBColor(0x3D, 0x8B, 0xD4)
AMBER_HEX = "F2A541"
AMBER = RGBColor(0xF2, 0xA5, 0x41)
LIGHT_BLUE = RGBColor(0x9C, 0xC6, 0xEE)
SIDE_TEXT = RGBColor(0xE8, 0xEE, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1E, 0x25, 0x30)
MUTED = RGBColor(0x5B, 0x66, 0x75)

PAGE_W, PAGE_H = 8.27, 11.69          # A4
PAGE_MARGIN_TB = 0.35
SIDE_W = 2.55
MAIN_W = PAGE_W - SIDE_W
CELL_PAD_X, CELL_PAD_Y = 0.28, 0.2
# Google Docs (and Word in compatibility mode) places a table's *text* at the page
# margin, so with zero page margins the cell padding hangs off the left edge.
# Indent the table by the padding and shrink the main column to compensate.
MAIN_W -= CELL_PAD_X
MAIN_TEXT_W = MAIN_W - 2 * CELL_PAD_X

CONTACT_LABELS = {"location": "Location", "phone": "Phone", "email": "Email", "linkedin": "LinkedIn"}


# ---- helpers ---------------------------------------------------------------

def make_round_photo(size=600, ring=14):
    """Circle-crop the photo with a white ring on a transparent background."""
    src = Image.open(PHOTO).convert("RGB")
    side = min(src.size)
    top = min(int(src.height * 0.03), src.height - side)  # portrait: keep the head, trim the chest
    src = src.crop(((src.width - side) // 2, top, (src.width + side) // 2, top + side))
    src = src.resize((size, size), Image.LANCZOS)
    big = size * 4  # supersample for smooth edges
    mask = Image.new("L", (big, big), 0)
    ImageDraw.Draw(mask).ellipse((0, 0, big - 1, big - 1), fill=255)
    mask = mask.resize((size, size), Image.LANCZOS)
    inner = Image.new("L", (big, big), 0)
    ImageDraw.Draw(inner).ellipse((ring * 4, ring * 4, big - 1 - ring * 4, big - 1 - ring * 4), fill=255)
    inner = inner.resize((size, size), Image.LANCZOS)
    out = Image.new("RGBA", (size, size), (255, 255, 255, 0))
    out.paste(Image.new("RGBA", (size, size), (255, 255, 255, 255)), (0, 0), mask)  # white ring
    out.paste(src, (0, 0), inner)
    out.save(ROUND_PHOTO)


def run(p, text, size=9.5, bold=False, italic=False, color=INK, caps=False):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.bold, r.italic, r.font.all_caps = bold, italic, caps
    r.font.color.rgb = color
    return r


def para(cell, before=0, after=2, align=None, line=1.15):
    p = cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    if align is not None:
        p.alignment = align
    return p


# Child-element order required by the Word schema. Word tolerates out-of-order
# properties; Google Docs silently drops them (layout, padding, shading).
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
               "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
               "shd", "tblLayout", "tblCellMar", "tblLook"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
              "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
             "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
             "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
             "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
             "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
             "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]


def put(parent, el, order):
    """Insert el into parent at its schema position, replacing any existing element of that tag."""
    name = el.tag.split("}")[1]
    for old in parent.findall(qn(f"w:{name}")):
        parent.remove(old)
    rank = order.index(name)
    for child in parent:
        child_name = child.tag.split("}")[1]
        if child_name in order and order.index(child_name) > rank:
            child.addprevious(el)
            return
    parent.append(el)


def margins_el(tag, pad_x, pad_y):
    mar = OxmlElement(tag)
    for side, val in (("top", pad_y), ("left", pad_x), ("bottom", pad_y), ("right", pad_x)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(val * 1440)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    return mar


def bottom_border(p, color_hex, size=8):
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in {"val": "single", "sz": str(size), "space": "2", "color": color_hex}.items():
        b.set(qn(f"w:{k}"), v)
    pbdr.append(b)
    put(p._p.get_or_add_pPr(), pbdr, PPR_ORDER)


def shade(cell, fill_hex):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    put(cell._tc.get_or_add_tcPr(), shd, TCPR_ORDER)


def setup_table(table, widths_in, pad_x, pad_y, border=None, indent=0):
    """Fixed widths, cell padding, and either no borders or (size, color) borders."""
    tblPr = table._tbl.tblPr
    tblInd = OxmlElement("w:tblInd")
    tblInd.set(qn("w:w"), str(int(indent * 1440)))
    tblInd.set(qn("w:type"), "dxa")
    put(tblPr, tblInd, TBLPR_ORDER)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(int(sum(widths_in) * 1440)))
    tblW.set(qn("w:type"), "dxa")
    put(tblPr, tblW, TBLPR_ORDER)

    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        if border:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(border[0]))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), border[1])
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    put(tblPr, borders, TBLPR_ORDER)

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    put(tblPr, layout, TBLPR_ORDER)
    put(tblPr, margins_el("w:tblCellMar", pad_x, pad_y), TBLPR_ORDER)

    for col, w in zip(table.columns, widths_in):
        col.width = Inches(w)
    for row in table.rows:
        for c, w in zip(row.cells, widths_in):
            c.width = Inches(w)
            # Per-cell padding too, for importers that ignore the table default.
            put(c._tc.get_or_add_tcPr(), margins_el("w:tcMar", pad_x, pad_y), TCPR_ORDER)


def side_heading(cell, text):
    p = para(cell, before=12, after=6)
    run(p, text, size=10.5, bold=True, color=WHITE, caps=True)
    bottom_border(p, AMBER_HEX, size=12)


def main_heading(cell, text):
    p = para(cell, before=10, after=6)
    p.paragraph_format.keep_with_next = True
    run(p, text, size=11, bold=True, color=NAVY, caps=True)
    bottom_border(p, AMBER_HEX, size=14)


def bullet(cell, text):
    p = cell.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(1.5), 1.1
    lead, rest = split_lead(text)
    if lead:
        run(p, lead, bold=True, color=NAVY)
        run(p, " — " + rest)
    else:
        run(p, rest)


# ---- build -----------------------------------------------------------------

def build_sidebar(cell):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    p.add_run().add_picture(str(ROUND_PHOTO), width=Inches(1.6), height=Inches(1.6))

    side_heading(cell, "Contact")
    for key, value in CONTACT.items():
        if value:
            run(para(cell, after=0), CONTACT_LABELS[key], size=8.5, bold=True, color=LIGHT_BLUE, caps=True)
            run(para(cell, after=4), value, size=9, color=SIDE_TEXT)

    side_heading(cell, "Certifications")
    for c in CERTIFICATIONS:
        p = para(cell, after=5)
        run(p, "✔ ", size=9, bold=True, color=AMBER)
        run(p, c, size=9, color=SIDE_TEXT)

    side_heading(cell, "Education")
    for degree, school, year in EDUCATION:
        run(para(cell, after=1), degree, size=9, bold=True, color=WHITE)
        run(para(cell, after=1), school, size=9, color=SIDE_TEXT)
        run(para(cell, after=4), year, size=9, color=SIDE_TEXT)

    side_heading(cell, "Skills")
    for label, items in SIDEBAR_SKILLS:
        run(para(cell, after=0), label, size=8.5, bold=True, color=LIGHT_BLUE, caps=True)
        run(para(cell, after=5), items, size=9, color=SIDE_TEXT)

    side_heading(cell, "Expertise")
    run(para(cell, after=0), " · ".join(CORE_COMPETENCIES), size=9, color=SIDE_TEXT)


def build_stats(cell):
    n = len(HIGHLIGHTS)
    table = cell.add_table(rows=1, cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Thick white borders read as gaps between the shaded boxes.
    setup_table(table, [MAIN_TEXT_W / n] * n, pad_x=0.05, pad_y=0.06, border=(24, "FFFFFF"))
    for c, (num, label) in zip(table.rows[0].cells, HIGHLIGHTS):
        shade(c, NAVY_HEX)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run(p, num, size=17, bold=True, color=AMBER)
        p = c.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run(p, label, size=8, color=SIDE_TEXT)


def build_main(cell):
    first, _, last = NAME.partition(" ")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run(p, first + " ", size=28, bold=True, color=NAVY)
    run(p, last, size=28, color=NAVY)
    p = para(cell, after=4)
    run(p, HEADLINE, size=10.5, color=MUTED)
    p = para(cell, before=2, after=4)
    run(p, TAGLINE, size=11, italic=True, color=ACCENT)
    bottom_border(p, NAVY_HEX, size=12)

    main_heading(cell, "Profile")
    run(para(cell, after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY), SUMMARY)
    build_stats(cell)

    main_heading(cell, "Work Experience")
    for job in EXPERIENCE:
        p = para(cell, before=7, after=0)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.tab_stops.add_tab_stop(Inches(MAIN_TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
        run(p, job["title"], size=10.5, bold=True, color=NAVY)
        run(p, "\t" + job["dates"], size=9, bold=True, color=MUTED)
        p = para(cell, after=2)
        p.paragraph_format.keep_with_next = True
        run(p, f'{job["company"]} · {job["location"]}', size=9.5, bold=True, color=ACCENT)
        for b in job["bullets"]:
            bullet(cell, b)


def build_docx():
    make_round_photo()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(9.5)

    s = doc.sections[0]
    s.page_width, s.page_height = Inches(PAGE_W), Inches(PAGE_H)
    s.left_margin = s.right_margin = Inches(0)
    s.top_margin = s.bottom_margin = Inches(PAGE_MARGIN_TB)
    s.header_distance = s.footer_distance = Inches(0.2)

    table = doc.add_table(rows=1, cols=2)
    setup_table(table, [SIDE_W, MAIN_W], pad_x=CELL_PAD_X, pad_y=CELL_PAD_Y, indent=CELL_PAD_X)
    side, main = table.rows[0].cells
    shade(side, NAVY_HEX)
    build_sidebar(side)
    build_main(main)

    # Word adds an empty paragraph after a final table; add a tiny one so it can't spill onto a new page.
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(0)
    tail.paragraph_format.line_spacing = Pt(1)
    mark = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "2")
    mark.append(sz)
    put(tail._p.get_or_add_pPr(), mark, PPR_ORDER)

    doc.core_properties.title = f"{NAME} – Resume"
    doc.core_properties.author = NAME
    doc.save(DOCX_OUT)


if __name__ == "__main__":
    build_docx()
    print(f"Wrote {DOCX_OUT.name}")
