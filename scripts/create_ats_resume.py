"""Generate an ATS-friendly resume as DOCX and PDF from resume.md.

ATS rules followed: single column, no tables/text boxes/images, contact info in
the body (not the page header), standard section headings, plain bullets,
consistent "Mon YYYY" dates, and common fonts.
"""
from pathlib import Path
import html
import subprocess

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent  # the resume folder (scripts/ lives inside it)
DOCX_OUT = ROOT / "Abhishek_Gupta_Resume_ATS.docx"
PDF_OUT = ROOT / "Abhishek_Gupta_Resume_ATS.pdf"
HTML_TMP = ROOT / "preview" / "ats_resume.html"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

# ---- Content (edit resume.md, not this file) --------------------------------

from resume_content import (  # noqa: E402
    CERTIFICATIONS, CONTACT, CORE_COMPETENCIES, EDUCATION, EXPERIENCE, HEADLINE, NAME, SKILLS, SUMMARY,
)


def contact_line():
    return " | ".join(v for v in CONTACT.values() if v)


# ---- DOCX ------------------------------------------------------------------

FONT = "Calibri"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)


def _run(p, text, size=10, bold=False, italic=False, color=None):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.bold, r.italic = bold, italic
    if color:
        r.font.color.rgb = color
    return r


def _para(doc, before=0, after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    return p


def _section(doc, title):
    p = _para(doc, before=10, after=4)
    _run(p, title.upper(), size=12, bold=True, color=ACCENT)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in {"val": "single", "sz": "8", "space": "1", "color": "1F3A5F"}.items():
        bottom.set(qn(f"w:{k}"), v)
    pbdr.append(bottom)
    # pBdr must precede spacing in pPr; Google Docs drops out-of-order properties.
    p._p.get_or_add_pPr().insert(0, pbdr)


def _bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    _run(p, text)


def build_docx():
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.5)
        s.page_height, s.page_width = Inches(11.69), Inches(8.27)
        s.left_margin = s.right_margin = Inches(0.6)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(10)

    p = _para(doc, after=0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, NAME, size=20, bold=True, color=ACCENT)
    p = _para(doc, after=0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, HEADLINE, size=11, bold=True)
    p = _para(doc, after=4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, contact_line(), size=10.5)

    _section(doc, "Professional Summary")
    _run(_para(doc), SUMMARY)

    _section(doc, "Core Competencies")
    _run(_para(doc), " | ".join(CORE_COMPETENCIES))

    _section(doc, "Technical Skills")
    for label, items in SKILLS:
        p = _para(doc, after=1)
        _run(p, f"{label}: ", bold=True)
        _run(p, items)

    _section(doc, "Certifications")
    for c in CERTIFICATIONS:
        _bullet(doc, c)

    _section(doc, "Education")
    for degree, school, year in EDUCATION:
        p = _para(doc, after=0)
        _run(p, degree, bold=True)
        _run(p, f" | {year}")
        _run(_para(doc), school, italic=True)

    _section(doc, "Professional Experience")
    for job in EXPERIENCE:
        p = _para(doc, before=6, after=0)
        p.paragraph_format.keep_with_next = True
        _run(p, job["title"], size=11, bold=True)
        _run(p, f" | {job['dates']}", size=11)
        p = _para(doc, after=2)
        p.paragraph_format.keep_with_next = True
        _run(p, f"{job['company']} | {job['location']}", italic=True)
        for b in job["bullets"]:
            _bullet(doc, b)

    doc.core_properties.title = f"{NAME} – Resume"
    doc.core_properties.author = NAME
    doc.save(DOCX_OUT)


# ---- PDF (HTML rendered by headless Chrome, so text stays selectable) -------

def build_html():
    e = html.escape
    parts = [f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<title>{e(NAME)} – Resume</title><style>
@page {{ size: A4; margin: 0.5in 0.6in; }}
body {{ font-family: Calibri, Carlito, Arial, Helvetica, sans-serif; font-size: 10pt;
       line-height: 1.25; color: #111; margin: 0; }}
h1 {{ font-size: 20pt; color: #1F3A5F; text-align: center; margin: 0; }}
.headline {{ text-align: center; font-weight: bold; font-size: 11pt; margin: 2px 0; }}
.contact {{ text-align: center; margin: 0 0 6px; }}
h2 {{ font-size: 12pt; color: #1F3A5F; text-transform: uppercase; border-bottom: 1.2px solid #1F3A5F;
      margin: 9px 0 4px; padding-bottom: 1px; break-after: avoid; }}
p {{ margin: 0 0 3px; }}
.job {{ break-inside: avoid-page; margin-top: 5px; }}
.job h3 {{ font-size: 11pt; margin: 0; }}
.job h3 span {{ font-weight: normal; }}
.org {{ font-style: italic; margin: 0 0 2px; }}
ul {{ margin: 0; padding-left: 18px; }}
li {{ margin: 0 0 1px; }}
</style></head><body>
<h1>{e(NAME)}</h1>
<p class="headline">{e(HEADLINE)}</p>
<p class="contact">{e(contact_line())}</p>
<h2>Professional Summary</h2><p>{e(SUMMARY)}</p>
<h2>Core Competencies</h2><p>{e(' | '.join(CORE_COMPETENCIES))}</p>
<h2>Technical Skills</h2>"""]
    parts += [f"<p><b>{e(l)}:</b> {e(i)}</p>" for l, i in SKILLS]
    parts.append("<h2>Certifications</h2><ul>" + "".join(f"<li>{e(c)}</li>" for c in CERTIFICATIONS) + "</ul>")
    parts.append("<h2>Education</h2>")
    for degree, school, year in EDUCATION:
        parts.append(f"<p><b>{e(degree)}</b> | {e(year)}</p><p><i>{e(school)}</i></p>")
    parts.append("<h2>Professional Experience</h2>")
    for j in EXPERIENCE:
        items = "".join(f"<li>{e(b)}</li>" for b in j["bullets"])
        parts.append(
            f'<div class="job"><h3>{e(j["title"])} <span>| {e(j["dates"])}</span></h3>'
            f'<p class="org">{e(j["company"])} | {e(j["location"])}</p><ul>{items}</ul></div>'
        )
    parts.append("</body></html>")
    return "\n".join(parts)


def build_pdf():
    HTML_TMP.parent.mkdir(exist_ok=True)
    HTML_TMP.write_text(build_html(), encoding="utf-8")
    subprocess.run(
        [CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
         f"--print-to-pdf={PDF_OUT}", HTML_TMP.as_uri()],
        check=True, capture_output=True,
    )


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(f"Wrote {DOCX_OUT.name} and {PDF_OUT.name}")
